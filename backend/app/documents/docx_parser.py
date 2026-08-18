import logging
from pathlib import Path
import zipfile
from backend.app.core.config import settings
from backend.app.documents.sanitizer import DocumentSanitizer

logger = logging.getLogger(__name__)


class DOCXParser:
    """Extracts native paragraph and table text from Microsoft Word .docx documents with ZIP bomb security guardrails."""

    @classmethod
    def parse_docx(cls, file_path: Path | str) -> tuple[str, int, str, bool]:
        """
        Parse .docx file and return (extracted_text, page_count=1, extraction_method="NATIVE_TEXT", is_ocr_applied=False).
        
        Security Invariants:
            1. Enforces ZIP decompression size and ratio checks before opening (ZIP bomb defense).
            2. VBA Macros, embedded executables, and external XML links are completely IGNORED and NEVER executed.
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"DOCX file not found at {path}")

        # 1. ZIP Bomb Decompression Security Check
        cls._verify_zip_security(path)

        # 2. Extract Text via python-docx
        try:
            import docx
        except ImportError as exc:
            raise RuntimeError("python-docx package is not installed") from exc

        try:
            doc = docx.Document(str(path))
            text_blocks: list[str] = []

            # Extract paragraph text
            for para in doc.paragraphs:
                clean_para = DocumentSanitizer.sanitize_text(para.text)
                if clean_para:
                    text_blocks.append(clean_para)

            # Extract table cell contents
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [DocumentSanitizer.sanitize_text(cell.text) for cell in row.cells]
                    clean_row = " | ".join([c for c in row_cells if c])
                    if clean_row:
                        text_blocks.append(clean_row)

            full_docx_text = "\n".join(text_blocks)
            formatted_text, _ = DocumentSanitizer.format_page_text([full_docx_text])
            return formatted_text, 1, "NATIVE_TEXT", False

        except Exception as exc:
            raise RuntimeError(f"DOCX parsing failed: {exc}") from exc

    @classmethod
    def _verify_zip_security(cls, path: Path) -> None:
        """Inspect DOCX ZIP archive compression metrics to prevent decompression bomb attacks."""
        if not zipfile.is_zipfile(path):
            raise ValueError(f"File {path.name} is not a valid DOCX/ZIP archive")

        total_uncompressed = 0
        total_compressed = 0

        with zipfile.ZipFile(path, "r") as zf:
            for info in zf.infolist():
                total_uncompressed += info.file_size
                total_compressed += info.compress_size

                # Block executable file extensions hidden inside archive
                lower_name = info.filename.lower()
                if lower_name.endswith((".exe", ".dll", ".bat", ".cmd", ".vbs", ".ps1", ".jar")):
                    raise ValueError(f"Malicious executable entry detected in DOCX archive: {info.filename}")

        # Check total uncompressed size limit
        if total_uncompressed > settings.DOC_ZIP_MAX_UNCOMPRESSED_BYTES:
            raise ValueError(
                f"DOCX uncompressed size ({total_uncompressed} bytes) exceeds limit ({settings.DOC_ZIP_MAX_UNCOMPRESSED_BYTES} bytes)"
            )

        # Check compression ratio limit
        if total_compressed > 0:
            ratio = total_uncompressed / total_compressed
            if ratio > settings.DOC_ZIP_MAX_RATIO:
                raise ValueError(
                    f"DOCX ZIP compression ratio ({ratio:.1f}:1) exceeds safety threshold ({settings.DOC_ZIP_MAX_RATIO}:1)"
                )
